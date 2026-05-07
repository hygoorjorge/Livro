"""Generate a DOCX rendition of the (edited) book.

Reconstructs paragraphs from SQLite. The DOCX is best-effort — fonts/spacing
approximate the original PDF but layout fidelity is not guaranteed (the
authoritative output is the in-place edited PDF).
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.db.models import Page, Paragraph, ParagraphKind, TextSpan


def _color_to_rgb(color: int) -> RGBColor:
    return RGBColor((color >> 16) & 0xFF, (color >> 8) & 0xFF, color & 0xFF)


async def export_docx(db: AsyncSession, session_id: int, output_path: str | Path) -> Path:
    pages = (
        await db.execute(
            select(Page)
            .where(Page.session_id == session_id)
            .order_by(Page.page_num)
        )
    ).scalars().all()

    doc = Document()
    for page in pages:
        paragraphs = (
            await db.execute(
                select(Paragraph)
                .where(Paragraph.page_id == page.id)
                .order_by(Paragraph.paragraph_index)
            )
        ).scalars().all()
        for para in paragraphs:
            spans = (
                await db.execute(
                    select(TextSpan)
                    .where(TextSpan.page_id == page.id)
                    .order_by(TextSpan.span_index)
                )
            ).scalars().all()
            spans_by_idx = {s.span_index: s for s in spans}

            doc_para = doc.add_paragraph()
            if para.kind == ParagraphKind.block_quote.value:
                doc_para.paragraph_format.left_indent = Pt(28)
            elif para.kind == ParagraphKind.footnote.value:
                doc_para.paragraph_format.first_line_indent = Pt(0)

            if not para.span_ids_json:
                run = doc_para.add_run(para.full_text)
                continue

            for span_idx in para.span_ids_json:
                span = spans_by_idx.get(span_idx)
                if span is None:
                    continue
                run = doc_para.add_run(span.text + " ")
                run.font.size = Pt(span.font_size)
                run.font.color.rgb = _color_to_rgb(span.color)
                fname = (span.font_name or "").lower()
                if "bold" in fname:
                    run.bold = True
                if "italic" in fname or "oblique" in fname:
                    run.italic = True
        doc.add_page_break()

    output_path = Path(output_path)
    doc.save(str(output_path))
    return output_path
